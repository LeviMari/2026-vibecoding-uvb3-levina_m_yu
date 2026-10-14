"""Выгрузка задач и встреч в Excel."""

from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from dates import parse_date, today
from storage import PRIORITY_RU, STATUS_RU


class ExportError(Exception):
    def __init__(self, user_message: str) -> None:
        super().__init__(user_message)
        self.user_message = user_message


def _fmt_date(value: str | date | None) -> str:
    d = parse_date(value) if isinstance(value, str) else value
    if d is None:
        # completed_at может быть datetime ISO
        if isinstance(value, str) and len(value) >= 10:
            d = parse_date(value[:10])
        if d is None:
            return ""
    return d.strftime("%d.%m.%Y")


def _autosize(ws) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = 10
        for cell in col:
            text = str(cell.value) if cell.value is not None else ""
            width = max(width, min(len(text) + 2, 60))
        ws.column_dimensions[letter].width = width


def _write_header(ws, titles: list[str]) -> None:
    bold = Font(bold=True)
    for i, title in enumerate(titles, start=1):
        cell = ws.cell(1, i, title)
        cell.font = bold


def build_tasks_xlsx(user: dict[str, Any]) -> bytes:
    """Листы: открытые задачи, выполненные за 30 дней, встречи."""
    try:
        wb = Workbook()
        open_ws = wb.active
        open_ws.title = "Открытые задачи"
        _write_header(open_ws, ["Название", "Срок", "Приоритет", "Статус", "Создана"])
        row = 2
        for t in user.get("tasks") or []:
            if t.get("status") == "done":
                continue
            open_ws.cell(row, 1, t.get("title") or "")
            open_ws.cell(row, 2, _fmt_date(t.get("deadline")))
            open_ws.cell(row, 3, PRIORITY_RU.get(t.get("priority"), t.get("priority") or ""))
            open_ws.cell(row, 4, STATUS_RU.get(t.get("status"), t.get("status") or ""))
            open_ws.cell(row, 5, _fmt_date(t.get("created_at")))
            row += 1
        _autosize(open_ws)

        done_ws = wb.create_sheet("Выполненные")
        _write_header(done_ws, ["Название", "Срок", "Приоритет", "Закрыта"])
        cutoff = (today() - timedelta(days=30)).isoformat()
        row = 2
        for t in user.get("tasks") or []:
            if t.get("status") != "done":
                continue
            completed = str(t.get("completed_at") or "")[:10]
            if completed and completed < cutoff:
                continue
            done_ws.cell(row, 1, t.get("title") or "")
            done_ws.cell(row, 2, _fmt_date(t.get("deadline")))
            done_ws.cell(row, 3, PRIORITY_RU.get(t.get("priority"), t.get("priority") or ""))
            done_ws.cell(row, 4, _fmt_date(t.get("completed_at")))
            row += 1
        _autosize(done_ws)

        meet_ws = wb.create_sheet("Встречи")
        _write_header(meet_ws, ["Дата", "Тема", "Решения"])
        row = 2
        meetings = sorted(user.get("meetings") or [], key=lambda m: m.get("date") or "", reverse=True)
        for m in meetings:
            decisions = m.get("decisions") or []
            text = "\n".join(str(x) for x in decisions)
            meet_ws.cell(row, 1, _fmt_date(m.get("date")))
            meet_ws.cell(row, 2, m.get("title") or "")
            meet_ws.cell(row, 3, text).alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
        _autosize(meet_ws)
        meet_ws.column_dimensions["C"].width = 80  # решения многострочные — переносим текст

        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()
    except ExportError:
        raise
    except Exception as exc:
        raise ExportError("Не получилось собрать Excel. Попробуй ещё раз.") from exc


def export_filename(ref: date | None = None) -> str:
    d = ref or today()
    return f"tasks_{d.strftime('%d.%m.%Y')}.xlsx"


SURVEY_TITLES = {
    "liked": "Что понравилось",
    "disliked": "Что не понравилось",
    "improve": "Что улучшить",
    "would_use": "Будет пользоваться",
}


def build_admin_xlsx(data: dict) -> bytes:
    """Выгрузка для отчёта: пользователи, ответы на опрос, отзывы, действия по дням."""
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Пользователи"
        _write_header(ws, ["Имя", "Действий", "Последняя активность", "Встреч", "Задач", "Календарь", "Опрос пройден"])
        for row, u in enumerate(data.get("users") or [], start=2):
            ws.cell(row, 1, u.get("name") or "")
            ws.cell(row, 2, u.get("actions") or 0)
            ws.cell(row, 3, (u.get("last_active") or "")[:16].replace("T", " "))
            ws.cell(row, 4, u.get("meetings") or 0)
            ws.cell(row, 5, u.get("tasks") or 0)
            ws.cell(row, 6, "да" if u.get("has_calendar") else "нет")
            ws.cell(row, 7, _fmt_date(u.get("survey_done_at")) or "нет")
        _autosize(ws)

        ws = wb.create_sheet("Опрос")
        _write_header(ws, ["Дата", "Имя", "Вопрос", "Ответ"])
        for row, a in enumerate(data.get("survey") or [], start=2):
            ws.cell(row, 1, _fmt_date(a.get("created_at")))
            ws.cell(row, 2, a.get("name") or "")
            ws.cell(row, 3, SURVEY_TITLES.get(a.get("question"), a.get("question")))
            ws.cell(row, 4, a.get("answer") or "").alignment = Alignment(wrap_text=True, vertical="top")
        _autosize(ws)
        ws.column_dimensions["D"].width = 80

        ws = wb.create_sheet("Отзывы")
        _write_header(ws, ["Дата", "Имя", "Отзыв"])
        for row, f in enumerate(data.get("feedback") or [], start=2):
            ws.cell(row, 1, _fmt_date(f.get("created_at")))
            ws.cell(row, 2, f.get("name") or "")
            ws.cell(row, 3, f.get("text") or "").alignment = Alignment(wrap_text=True, vertical="top")
        _autosize(ws)
        ws.column_dimensions["C"].width = 80

        ws = wb.create_sheet("Действия по дням")
        _write_header(ws, ["День", "Действие", "Сколько раз"])
        for row, d in enumerate(data.get("daily") or [], start=2):
            ws.cell(row, 1, _fmt_date(d.get("day")))
            ws.cell(row, 2, d.get("action") or "")
            ws.cell(row, 3, d.get("n") or 0)
        _autosize(ws)

        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()
    except Exception as exc:
        raise ExportError("Не получилось собрать выгрузку для администратора.") from exc
